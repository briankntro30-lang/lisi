import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

def export_excel(fig, dfs, state, kpis):

    fig.savefig("simogramme.png")

    with pd.ExcelWriter("simogramme.xlsx", engine="openpyxl") as writer:

        pd.concat(dfs).to_excel(writer, sheet_name="Data", index=False)

        wb = writer.book
        ws = wb.create_sheet("Report")

        ws["A1"] = "Cycle"
        ws["B1"] = kpis["temps_cycle"]

        img = XLImage("simogramme.png")
        ws.add_image(img, "D2")
